"""Parse and chunk files uploaded to the managed knowledge base."""

from __future__ import annotations

import hashlib
import io
import re
import zipfile
from dataclasses import dataclass
from pathlib import PurePosixPath
from typing import Optional

from docx import Document as WordDocument
from langchain.text_splitter import MarkdownHeaderTextSplitter, RecursiveCharacterTextSplitter
from langchain_core.documents import Document
from pypdf import PdfReader

from config import (
    CHUNK_OVERLAP,
    CHUNK_SIZE,
    HEADERS_TO_SPLIT_ON,
    TEXT_SPLITTER_SEPARATORS,
)

MAX_UPLOAD_BYTES = 4 * 1024 * 1024
MAX_SOURCE_CHUNKS = 500
MAX_EXTRACTED_CHARACTERS = 500_000
MAX_PDF_PAGES = 500
MAX_DOCX_UNCOMPRESSED_BYTES = 25 * 1024 * 1024
MAX_DOCX_ENTRIES = 10_000
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


@dataclass(frozen=True)
class ParsedSection:
    text: str
    page_number: Optional[int] = None


@dataclass(frozen=True)
class ParsedSource:
    filename: str
    extension: str
    sections: list[ParsedSection]
    content_hash: str

    @property
    def text(self) -> str:
        return "\n\n".join(section.text for section in self.sections)


def _safe_filename(filename: Optional[str]) -> str:
    name = PurePosixPath(str(filename or "").replace("\\", "/")).name.strip()
    if not name or name in {".", ".."}:
        raise ValueError("Choose a file with a valid filename.")
    if len(name) > 180 or any(ord(character) < 32 for character in name):
        raise ValueError("The filename is too long or contains invalid characters.")
    return name


def _check_text_limit(sections: list[ParsedSection]) -> None:
    if sum(len(section.text) for section in sections) > MAX_EXTRACTED_CHARACTERS:
        raise ValueError("This file contains too much extracted text. The limit is 500,000 characters.")
    if not any(section.text.strip() for section in sections):
        raise ValueError("No readable text was found in this file.")


def _parse_pdf(raw: bytes) -> list[ParsedSection]:
    if not raw.lstrip().startswith(b"%PDF"):
        raise ValueError("The file does not appear to be a valid PDF.")
    try:
        reader = PdfReader(io.BytesIO(raw), strict=False)
        if len(reader.pages) > MAX_PDF_PAGES:
            raise ValueError("PDF files are limited to 500 pages.")
        sections: list[ParsedSection] = []
        character_count = 0
        for page_number, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            character_count += len(text)
            if character_count > MAX_EXTRACTED_CHARACTERS:
                raise ValueError("This file contains too much extracted text. The limit is 500,000 characters.")
            if text:
                sections.append(ParsedSection(text=text, page_number=page_number))
        return sections
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("This PDF could not be read. Try exporting it as a text-based PDF.") from exc


def _parse_docx(raw: bytes) -> list[ParsedSection]:
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_DOCX_ENTRIES:
                raise ValueError("This Word file contains too many internal entries.")
            expanded_size = sum(entry.file_size for entry in entries)
            if expanded_size > MAX_DOCX_UNCOMPRESSED_BYTES:
                raise ValueError("Word files are limited to 25 MB after decompression.")
            if "word/document.xml" not in archive.namelist():
                raise ValueError("This file does not appear to be a valid DOCX document.")

        document = WordDocument(io.BytesIO(raw))
        blocks = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                row_text = " | ".join(value for value in cells if value)
                if row_text:
                    blocks.append(row_text)
        sections = [ParsedSection(text="\n\n".join(blocks))] if blocks else []
        _check_text_limit(sections)
        return sections
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("This Word document could not be read. Upload a valid DOCX file.") from exc


def parse_uploaded_file(filename: Optional[str], raw: bytes) -> ParsedSource:
    """Extract text from an allowed upload and return a content fingerprint."""
    safe_name = _safe_filename(filename)
    extension = PurePosixPath(safe_name.lower()).suffix
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError("Supported file types are PDF, DOCX, TXT, and MD.")
    if not raw:
        raise ValueError("The selected file is empty.")
    if len(raw) > MAX_UPLOAD_BYTES:
        raise ValueError("Files must be 4 MB or smaller.")

    if extension == ".pdf":
        sections = _parse_pdf(raw)
    elif extension == ".docx":
        sections = _parse_docx(raw)
    else:
        try:
            text = raw.decode("utf-8-sig").strip()
        except UnicodeDecodeError as exc:
            raise ValueError("Text files must use UTF-8 encoding.") from exc
        sections = [ParsedSection(text=text)] if text else []

    _check_text_limit(sections)
    normalized = re.sub(r"\s+", " ", " ".join(section.text for section in sections)).strip()
    content_hash = hashlib.sha256(normalized.encode("utf-8")).hexdigest()
    return ParsedSource(
        filename=safe_name,
        extension=extension,
        sections=sections,
        content_hash=content_hash,
    )


def chunk_uploaded_source(parsed: ParsedSource, scope: str) -> list[Document]:
    """Split extracted text while retaining filename, page, and Markdown headers."""
    recursive_splitter = RecursiveCharacterTextSplitter(
        chunk_size=CHUNK_SIZE,
        chunk_overlap=CHUNK_OVERLAP,
        length_function=len,
        separators=TEXT_SPLITTER_SEPARATORS,
    )
    chunks: list[Document] = []

    if parsed.extension == ".md":
        markdown_splitter = MarkdownHeaderTextSplitter(
            headers_to_split_on=HEADERS_TO_SPLIT_ON,
            strip_headers=False,
        )
        sections = markdown_splitter.split_text(parsed.text)
        for section in sections:
            section.metadata.update({"source": parsed.filename, "format": "md", "scope": scope})
            chunks.extend(
                recursive_splitter.split_documents([section])
                if len(section.page_content) > CHUNK_SIZE
                else [section]
            )
    else:
        for section in parsed.sections:
            metadata = {
                "source": parsed.filename,
                "format": parsed.extension.lstrip("."),
                "scope": scope,
            }
            if section.page_number is not None:
                metadata["page"] = section.page_number
            chunks.extend(recursive_splitter.create_documents([section.text], metadatas=[metadata]))

    if len(chunks) > MAX_SOURCE_CHUNKS:
        raise ValueError("This file creates more than the 500 chunk limit. Upload a smaller document.")
    if not chunks:
        raise ValueError("No searchable text chunks could be created from this file.")
    return chunks
